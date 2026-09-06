from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path

OUT = Path(r"C:\Users\thors\OneDrive\Documents\GitHub\impfgruppe-rgzv-hagen\Satzung Vorabentwurf Finanzamt NEU 31.08.docx")

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.90)
sec.right_margin = Inches(0.90)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Liberation Sans'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Liberation Sans')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Liberation Sans')
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5.5)
normal.paragraph_format.line_spacing = 1.0

if 'Satzung Heading' not in styles:
    hs = styles.add_style('Satzung Heading', WD_STYLE_TYPE.PARAGRAPH)
else:
    hs = styles['Satzung Heading']
hs.font.name = 'Calibri'
hs._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
hs._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
hs.font.size = Pt(13)
hs.font.bold = True
hs.font.color.rgb = RGBColor(54, 95, 145)
hs.paragraph_format.space_before = Pt(14)
hs.paragraph_format.space_after = Pt(2)
hs.paragraph_format.keep_with_next = True

def set_font(run, name='Liberation Sans', size=10.5, bold=False, italic=False, strike=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.strike = strike
    if color:
        run.font.color.rgb = RGBColor(*color)

def body(text='', after=5.5, before=0, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = keep
    set_font(p.add_run(text))
    return p

def mixed(parts, after=5.5, before=0, left=0, first=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.left_indent = Inches(left)
    p.paragraph_format.first_line_indent = Inches(first)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    for text, strike in parts:
        r = p.add_run(text); set_font(r, strike=strike)
    return p

def heading(text, before=None):
    p = doc.add_paragraph(style='Satzung Heading')
    if before is not None: p.paragraph_format.space_before = Pt(before)
    set_font(p.add_run(text), 'Calibri', 13, True, color=(54,95,145))
    return p

def item(text, after=5.5, strike=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.23)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    set_font(p.add_run(text), strike=strike)
    return p

def pagebreak():
    doc.add_page_break()

# Page 1
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(9)
set_font(p.add_run('Satzung'), size=18, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(7)
set_font(p.add_run('des Rassegeflügelzuchtvereins Hagen und Umgebung seit 1903 e.V.'), size=13, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(24)
set_font(p.add_run('Neufassung 2026 - Entwurf zur Vorabprüfung durch Finanzamt und Registergericht'), size=9.5, italic=True)

heading('§ 1 Name, Sitz, Eintragung und Geschäftsjahr', before=0)
body('(1) Der Verein führt den Namen „Rassegeflügelzuchtverein Hagen und Umgebung seit 1903 e.V.“.')
mixed([('(2) Der Verein hat seinen Sitz in Hagen und ist ', False), ('in das Vereinsregister eingetragen.', True)])
body('(3) Das Geschäftsjahr ist das Kalenderjahr.')

heading('§ 2 Verbandszugehörigkeit')
body('Der Verein kann Mitglied in fachlich zuständigen Kreis-, Landes- und Bundesverbänden der Rassegeflügelzucht sein. Über Beitritt und Austritt entscheidet der Vorstand, soweit die Mitgliederversammlung sich die Entscheidung nicht vorbehält.')

heading('§ 3 Gemeinnützige Zwecke')
body('(1) Der Verein verfolgt ausschließlich und unmittelbar gemeinnützige Zwecke im Sinne des Abschnitts „Steuerbegünstigte Zwecke“ der Abgabenordnung.')
body('(2) Zwecke des Vereins sind:')
item('1. die Förderung der Tierzucht,')
item('2. die Förderung des Tierschutzes,')
item('3. die Förderung der Jugendhilfe sowie')
item('4. die Förderung des öffentlichen Gesundheitswesens und der öffentlichen Gesundheitspflege, insbesondere die Verhütung und Bekämpfung von Tierseuchen.')
body('(3) Der Verein ist politisch, religiös und weltanschaulich neutral. Er tritt für ein respektvolles und diskriminierungsfreies Vereinsleben ein.')

heading('§ 4 Verwirklichung der Satzungszwecke')
body('Die Satzungszwecke werden insbesondere verwirklicht durch:')
item('1. fachliche Beratung sowie öffentlich zugängliche Informations- und Bildungsveranstaltungen über Rasse- und Ziergeflügelzucht, artgerechte Haltung, Tiergesundheit und Tierschutz;')
item('2. Erhaltung, Dokumentation und verantwortungsvolle Weiterentwicklung von Geflügelrassen und Farbenschlägen, insbesondere gefährdeter Rassen;')
item('3. Durchführung fachlich ausgerichteter Rasse- und Ziergeflügelausstellungen, Zuchtschauen und Tierbewertungen nach anerkannten Bestimmungen;')
item('4. Förderung der Kinder- und Jugendarbeit, insbesondere durch Heranführung junger Menschen an die Rassegeflügelzucht, Vermittlung von Kenntnissen über Tierzucht, artgerechte Tierhaltung, Tierschutz und Tiergesundheit, Durchführung von Jugendveranstaltungen und Bildungsangeboten sowie Förderung von Verantwortungsbewusstsein, Gemeinschaft und Naturverständnis;')
item('5. Öffentlichkeitsarbeit zur Vermittlung von Wissen über Rassegeflügel, verantwortungsvolle Zucht und artgerechte Haltung;', after=0)

# Page 2
pagebreak()
item('6. Zusammenarbeit mit Tierärzten, Behörden, Schulen, Fachverbänden und steuerbegünstigten Einrichtungen;')
item('7. Informations-, Vorsorge- und Unterstützungsmaßnahmen zur Verhütung und Bekämpfung von Tierseuchen, soweit diese Maßnahmen den Satzungszwecken und der Allgemeinheit dienen.')

heading('§ 5 Selbstlosigkeit und Mittelverwendung')
body('(1) Der Verein ist selbstlos tätig; er verfolgt nicht in erster Linie eigenwirtschaftliche Zwecke.')
body('(2) Mittel des Vereins dürfen nur für die satzungsmäßigen Zwecke verwendet werden. Die Mitglieder erhalten keine Zuwendungen aus Mitteln des Vereins.')
body('(3) Es darf keine Person durch Ausgaben, die dem Zweck des Vereins fremd sind, oder durch unverhältnismäßig hohe Vergütungen begünstigt werden.')

heading('§ 6 Arten und Erwerb der Mitgliedschaft')
body('(1) Der Verein hat ordentliche Mitglieder, Jugendmitglieder, fördernde Mitglieder und Ehrenmitglieder.')
body('(2) Ordentliches Mitglied kann jede natürliche Person ab Vollendung des 18. Lebensjahres werden. Jugendmitglied kann jede natürliche Person ab Vollendung des 6. und bis zur Vollendung des 18. Lebensjahres werden. Juristische Personen können fördernde Mitglieder ohne Stimmrecht werden.')
body('(3) Der Aufnahmeantrag ist in Textform an den Vorstand zu richten. Der Vorstand entscheidet. Gegen eine Ablehnung kann innerhalb eines Monats nach Zugang der Ablehnung die Mitgliederversammlung angerufen werden; diese entscheidet endgültig.')
body('(4) Ehrenmitglieder werden auf Vorschlag des Vorstands durch die Mitgliederversammlung ernannt.')

heading('§ 7 Beendigung der Mitgliedschaft')
body('(1) Die Mitgliedschaft endet durch Tod, bei juristischen Personen durch deren Auflösung, durch Austritt oder Ausschluss.')
body('(2) Der Austritt ist gegenüber dem Vorstand in Textform zu erklären und wird mit Zugang wirksam. Bereits fällige Beiträge bleiben geschuldet.')
body('(3) Ein Mitglied kann aus wichtigem Grund ausgeschlossen werden, insbesondere wenn es in erheblicher Weise gegen die Interessen oder die Satzung des Vereins verstößt. Vor der Entscheidung ist ihm Gelegenheit zur Stellungnahme zu geben. Der Vorstand entscheidet mit Mehrheit und teilt die Entscheidung mit Gründen in Textform mit. Innerhalb eines Monats nach Zugang kann das Mitglied die Mitgliederversammlung anrufen; bis zu deren Entscheidung ruhen die Mitgliedschaftsrechte.')

heading('§ 8 Rechte und Pflichten')
body('(1) Stimmberechtigt sind ordentliche Mitglieder ab Vollendung des 18. Lebensjahres. Jedes stimmberechtigte Mitglied hat eine Stimme.')
body('(2) Die Mitglieder unterstützen die Satzungszwecke, beachten Satzung und wirksame Beschlüsse und gehen verantwortungsvoll mit Vereinseigentum um.')
body('(3) Ein Anspruch auf Leistungen oder Vermögen des Vereins besteht nicht.', after=0)

# Page 3
pagebreak()
heading('§ 9 Beiträge, Umlagen und Leistungsentgelte', before=0)
body('(1) Der Verein erhebt Mitgliedsbeiträge. Höhe, Fälligkeit und mögliche Ermäßigungen beschließt die Mitgliederversammlung in einer Beitragsordnung.')
body('(2) Zur Deckung eines außergewöhnlichen Finanzbedarfs kann die Mitgliederversammlung eine Umlage beschließen. Der Beschluss muss Zweck, Höhe und Fälligkeit bestimmen. Die Umlage darf je Mitglied innerhalb eines Geschäftsjahres das Zweifache des Jahresbeitrags nicht übersteigen.')
body('(3) Für besondere, individuell zurechenbare Leistungen und Veranstaltungen kann der Verein gesonderte Entgelte verlangen. Diese sind keine Mitgliedsbeiträge und werden getrennt beschlossen, abgerechnet und gebucht.')
heading('§ 10 Organe des Vereins')
body('Organe des Vereins sind die Mitgliederversammlung und der Vorstand.')
heading('§ 11 Mitgliederversammlung: Einberufung und Form')
body('(1) Mindestens einmal jährlich findet eine ordentliche Mitgliederversammlung statt.')
mixed([('(2) Sie wird vom ',False),('Vorsitz, bei Verhinderung vom stellvertretenden Vorsitz,',True),(' mit einer Frist von zwei Wochen in Textform unter Angabe der Tagesordnung einberufen. Für die Frist ist die Absendung an die zuletzt mitgeteilte Kontaktadresse maßgeblich.',False)])
body('(3) Die Mitgliederversammlung findet grundsätzlich als Präsenzversammlung statt. Bei der Einberufung kann vorgesehen werden, dass Mitglieder auch ohne Anwesenheit am Versammlungsort im Wege elektronischer Kommunikation teilnehmen und ihre Mitgliederrechte ausüben können (hybride Versammlung). Rein virtuelle Mitgliederversammlungen können einberufen werden, wenn die Mitgliederversammlung dies für künftige Versammlungen beschlossen hat. Bei einer hybriden oder virtuellen Versammlung ist in der Einladung anzugeben, wie die Mitglieder ihre Rechte im Wege elektronischer Kommunikation ausüben können.')
body('(4) Eine außerordentliche Mitgliederversammlung ist einzuberufen, wenn der Vorstand dies beschließt oder ein Drittel der stimmberechtigten Mitglieder dies unter Angabe von Zweck und Gründen verlangt.')
heading('§ 12 Aufgaben der Mitgliederversammlung')
body('Die Mitgliederversammlung ist insbesondere zuständig für:')
item('1. Wahl und Abberufung der Mitglieder des Vorstands;')
mixed([('2. Wahl von zwei ',False),('Kassenprüfenden,',True),(' die nicht dem Vorstand angehören dürfen;',False)], left=.23, first=-.14)
item('3. Entgegennahme der Jahres-, Tätigkeits- und Finanzberichte;')
item('4. Entgegennahme des Berichts der Kassenprüfung und Entlastung des Vorstands;')
item('5. Beschluss über Beitragsordnung, Umlagen und eine mögliche Vorstandsvergütung;')
item('6. Ernennung von Ehrenmitgliedern;')
item('7. Entscheidung über Berufungen gegen Aufnahmeablehnung und Ausschluss;')
item('8. Satzungsänderungen, Zweckänderungen und Auflösung des Vereins.')
heading('§ 13 Beschlüsse und Protokoll')
body('(1) Die ordnungsgemäß einberufene Mitgliederversammlung ist unabhängig von der Zahl der anwesenden oder wirksam zugeschalteten stimmberechtigten Mitglieder beschlussfähig.', after=0)

# Page 4
pagebreak()
body('(2) Soweit Gesetz oder Satzung nichts anderes bestimmen, werden Beschlüsse mit einfacher Mehrheit der abgegebenen gültigen Stimmen gefasst. Stimmenthaltungen gelten als nicht abgegeben. Bei Stimmengleichheit ist ein Antrag abgelehnt.')
body('(3) Satzungsänderungen bedürfen einer Mehrheit von drei Vierteln der abgegebenen gültigen Stimmen. Eine Änderung des Vereinszwecks richtet sich nach den gesetzlichen Vorschriften.')
body('(4) Wahlen erfolgen offen, sofern nicht ein stimmberechtigtes Mitglied geheime Wahl verlangt.')
body('(5) Über die Versammlung ist ein Ergebnisprotokoll zu erstellen, das von der Versammlungsleitung und der Protokollführung zu unterzeichnen ist.')
heading('§ 14 Vorstand nach § 26 BGB und Vertretung')
body('(1) Der Vorstand im Sinne des § 26 BGB besteht aus:')
item('1. dem Vorsitz,', strike=True)
item('2. dem stellvertretenden Vorsitz,', strike=True)
item('3. der Kassenführung und', strike=True)
item('4. der Schriftführung.', strike=True)
body('(2) Der Verein wird gerichtlich und außergerichtlich durch jeweils zwei Mitglieder des Vorstands gemeinsam vertreten.')
body('(3) Die Amtszeit beträgt drei Jahre. Wiederwahl ist zulässig. Der Vorstand bleibt bis zur Wahl eines neuen Vorstands im Amt.')
body('(4) Scheidet ein Vorstandsmitglied vorzeitig aus, kann der verbleibende Vorstand bis zur nächsten Mitgliederversammlung eine kommissarische Person aus dem Kreis der Vereinsmitglieder bestellen.')
heading('§ 15 Geschäftsführung und Beschlüsse des Vorstands')
body('(1) Der Vorstand führt die laufenden Geschäfte und ist für alle Angelegenheiten zuständig, die nicht einem anderen Organ zugewiesen sind. Er stellt insbesondere ordnungsgemäße Buchführung, Mittelverwendung, steuerliche Pflichten, Tätigkeitsnachweise und die Umsetzung der Mitgliederversammlungsbeschlüsse sicher.')
mixed([('(2) Vorstandssitzungen werden ',False),('vom Vorsitz, bei Verhinderung vom stellvertretenden Vorsitz, mit angemessener Frist einberufen.',True),(' Sitzungen können in Präsenz, telefonisch oder per Videokonferenz stattfinden.',False)])
mixed([('(3) Der Vorstand ist beschlussfähig, wenn mindestens drei Mitglieder teilnehmen. Er beschließt mit einfacher Mehrheit; bei Stimmengleichheit ',False),('ist der Antrag abgelehnt.',True),(' Beschlüsse in Textform sind zulässig, wenn kein Vorstandsmitglied widerspricht.',False)])
body('(4) Über Beschlüsse ist ein Protokoll zu führen. Bei einem persönlichen oder wirtschaftlichen Interessenkonflikt nimmt das betroffene Vorstandsmitglied an Beratung und Abstimmung nicht teil.')
heading('§ 16 Ehrenamt, Auslagen und Vergütung')
body('(1) Vereins- und Organämter werden grundsätzlich ehrenamtlich ausgeübt.')
body('(2) Nachgewiesene angemessene Auslagen können auf Grundlage eines Vorstandsbeschlusses erstattet werden.', after=0)

# Page 5
pagebreak()
body('(3) Die Mitgliederversammlung kann im Rahmen der finanziellen Möglichkeiten und der steuerrechtlich zulässigen Grenzen eine angemessene Vergütung für Vorstands- oder andere Vereinstätigkeiten beschließen. Betroffene Personen sind bei der Beschlussfassung nicht stimmberechtigt. Aufgaben, Umfang und Vergütung sind schriftlich zu vereinbaren.')
heading('§ 17 Erweiterter Vorstand und Beauftragte')
mixed([('(1) Zur Unterstützung können insbesondere ',False),('eine zweite Kassenführung, eine zweite Schriftführung,',True),(' Jugendleitung, Ausstellungsleitung, Zuchtwarte für Hühner und Tauben sowie ein Käfigwart gewählt oder bestellt werden.',False)])
body('(2) Der erweiterte Vorstand ist kein Vorstand im Sinne des § 26 BGB und nicht zur Vertretung des Vereins berechtigt. Aufgaben und Berichtspflichten legt der Vorstand in einer Geschäftsordnung fest.')
body('(3) Der Vorstand kann Beauftragte und Ausschüsse einsetzen, insbesondere für Gemeinnützigkeit, Tiergesundheit, Datenschutz oder Veranstaltungen. Die Verantwortung des Vorstands bleibt unberührt.')
heading('§ 18 Jugendgruppe')
body('(1) Die Jugendgruppe gestaltet ihre fachliche und soziale Jugendarbeit im Rahmen dieser Satzung und einer Jugendordnung.')
body('(2) Die Jugendleitung wird von den Jugendmitgliedern vorgeschlagen und von der Mitgliederversammlung bestätigt. Sie berichtet dem Vorstand und der Mitgliederversammlung.')
heading('§ 19 Kassenprüfung')
mixed([('(1) Die Mitgliederversammlung wählt ',False),('zwei Kassenprüfende für drei Jahre. Wiederwahl ist zulässig.',True)])
body('(2) Sie prüfen mindestens einmal jährlich die Ordnungsmäßigkeit der Buchführung, Bank- und Kassenbestände, Belege, Beschlüsse und die nachvollziehbare Zuordnung der Einnahmen und Ausgaben. Sie berichten der Mitgliederversammlung.')
heading('§ 20 Datenschutz, Vereinsordnungen und behördlich verlangte Satzungsanpassungen')
body('(1) Der Verein verarbeitet personenbezogene Daten nur, soweit dies für Mitgliedschaft, Vereinszwecke, Verbandsmeldungen und gesetzliche Pflichten erforderlich ist. Einzelheiten können in einer Datenschutzordnung geregelt werden.')
body('(2) Die Mitgliederversammlung kann Vereinsordnungen beschließen. Der Vorstand kann eine Geschäfts- und Finanzordnung erlassen, soweit sie dieser Satzung und Beschlüssen der Mitgliederversammlung nicht widerspricht.')
body('(3) Der Vorstand ist ermächtigt, Änderungen dieser Satzung zu beschließen, die das zuständige Finanzamt oder Registergericht für die Anerkennung der Gemeinnützigkeit oder die Eintragung verlangt, sofern dadurch die Vereinszwecke, die Mitgliedsrechte und die grundlegende Organisationsstruktur nicht wesentlich verändert werden. Die Mitglieder sind über solche Änderungen unverzüglich zu informieren. Soweit gesetzlich eine Beschlussfassung der Mitgliederversammlung zwingend erforderlich ist, bleibt diese unberührt.')
heading('§ 21 Auflösung und Vermögensbindung')
body('(1) Die Auflösung des Vereins kann nur eine eigens hierzu einberufene Mitgliederversammlung mit einer Mehrheit von drei Vierteln der abgegebenen gültigen Stimmen beschließen.', after=0)

# Page 6
pagebreak()
body('(2) Sofern die Mitgliederversammlung nichts anderes beschließt, sind zwei Mitglieder des Vorstands gemeinsam vertretungsberechtigte Liquidatoren.')
body('(3) Bei Auflösung oder Aufhebung des Vereins oder bei Wegfall steuerbegünstigter Zwecke fällt das Vermögen des Vereins an eine juristische Person des öffentlichen Rechts oder eine andere steuerbegünstigte Körperschaft zwecks Verwendung für die Förderung der Tierzucht und des Tierschutzes.')
heading('§ 22 Inkrafttreten')
body('(1) Diese Satzung wurde von der Mitgliederversammlung am ____________________ in Hagen beschlossen.')
body('(2) Sie tritt mit Eintragung in das Vereinsregister in Kraft und ersetzt die bisherige Satzung.', after=15)
p=body('', after=18); set_font(p.add_run('Beschlossen in Hagen am: '), bold=True); set_font(p.add_run('______________________________'))
p=body('Unterschriften gemäß den Anforderungen des Registerverfahrens:', after=20); p.runs[0].bold=True
for i in range(4):
    body('__________________________________________', after=13)

doc.core_properties.title = 'Satzung des Rassegeflügelzuchtvereins Hagen und Umgebung seit 1903 e.V.'
doc.core_properties.subject = 'Neufassung 2026 - Entwurf zur Vorabprüfung'
doc.core_properties.author = 'Rassegeflügelzuchtverein Hagen und Umgebung seit 1903 e.V.'
doc.save(OUT)
print(OUT)
